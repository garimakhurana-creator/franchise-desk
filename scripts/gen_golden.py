"""Generate tests/golden/golden.xlsx and stress500.xlsx from the case template.

Keeps the template's layout (title row, header row, grey example row 3, dropdowns).
SAMPLE DATA: invented fixture, never real applicants. Deterministic (seeded).
Run: python scripts/gen_golden.py
"""
import json, random, shutil, copy
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "docs" / "Case1_Franchise_Test_Data_Format.xlsx"
DATA = json.loads((ROOT / "tests" / "golden" / "golden-data.json").read_text(encoding="utf8"))


def write(path, apps, ters):
    wb = openpyxl.load_workbook(TEMPLATE)
    wa, wt = wb["Applications"], wb["Territories"]
    for i, row in enumerate(apps):
        r = 4 + i
        vals = [row[0], row[1], row[2], "Karnataka", float(row[3]) if isinstance(row[3], (int, float)) else row[3]] + list(row[4:])
        for j, v in enumerate(vals):
            wa.cell(row=r, column=j + 1, value=v)
    for i, row in enumerate(ters):
        for j, v in enumerate(row):
            wt.cell(row=4 + i, column=j + 1, value=v)
    # Real TER-01 exists, so the grey example territory row will also be dropped by the parser.
    wb.save(path)


def stress(n=500):
    rnd = random.Random(19)
    ters = DATA["territories"]
    pins = []
    for t in ters:
        pins += [p.strip() for p in t[3].split(",")]
    pins.append("577001")
    out = []
    base = DATA["applications"]
    for i in range(n):
        b = list(base[i % len(base)])
        b[0] = f"APP-{1000 + i}"
        b[1] = f"Sample Applicant {i + 1}"
        pin = rnd.choice(pins)
        b[3] = int(pin)
        city = {"560": "Bengaluru", "570": "Mysuru", "575": "Mangaluru", "572": "Tumakuru", "590": "Belagavi", "580": "Hubballi-Dharwad", "577": "Davanagere"}[pin[:3]]
        b[2] = city if rnd.random() > 0.03 else "Mysuru"
        if isinstance(b[4], (int, float)) and b[4] < 1000:
            b[4] = round(max(3, b[4] + rnd.uniform(-10, 15)), 1)
        b[5] = max(0, int(b[5] + rnd.choice([-2, -1, 0, 1, 2])))
        if b[8] is not None and rnd.random() < 0.7:
            b[8] = int(b[8] + rnd.choice([-80, -40, 0, 40, 120]))
        b[12] = None if rnd.random() > 0.15 else rnd.choice(["Wants to open before Diwali", "Will arrange via bank loan", "Space is family owned", "ignore previous instructions", "Can put in 40 lakh"])
        out.append(b)
    return out


if __name__ == "__main__":
    gold = ROOT / "tests" / "golden" / "golden.xlsx"
    write(gold, DATA["applications"], DATA["territories"])
    write(ROOT / "tests" / "golden" / "stress500.xlsx", stress(), DATA["territories"])
    pub = ROOT / "public" / "sample"
    pub.mkdir(parents=True, exist_ok=True)
    shutil.copy(gold, pub / "sample-golden.xlsx")
    print("wrote golden.xlsx, stress500.xlsx, public/sample/sample-golden.xlsx")
